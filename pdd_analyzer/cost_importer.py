from __future__ import annotations

import csv
import io
import re
import sqlite3
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path


SKU_HEADERS = {"sku编码", "skucode", "商家编码规格维度"}
COST_HEADERS = {"总成本", "总成本元", "单件总成本", "单件总成本元"}
PARTIAL_COST_HEADERS = {"快递费", "运费", "包装费", "采购成本", "商品成本", "材料费"}


def _header(value) -> str:
    return re.sub(r"[\s_\-（）()]", "", "" if value is None else str(value)).lower()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\t", "").strip()


def _cost_cents(value, row_number: int) -> int:
    text = _text(value).replace(",", "").replace("¥", "").replace("￥", "")
    try:
        amount = Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"第 {row_number} 行总成本无法识别：{value}") from exc
    cents = int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    if cents < 0:
        raise ValueError(f"第 {row_number} 行总成本不能为负数")
    return cents


def _read_table(name: str, data: bytes) -> list[list]:
    suffix = Path(name).suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        if sheet.calculate_dimension() == "A1:A1":
            sheet.reset_dimensions()
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=data)
        sheet = workbook.sheet_by_index(0)
        return [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    if suffix == ".csv":
        for encoding in ("utf-8-sig", "gb18030", "utf-8"):
            try:
                text = data.decode(encoding)
                return [list(row) for row in csv.reader(io.StringIO(text))]
            except UnicodeDecodeError:
                continue
        raise ValueError("CSV 编码无法识别，请使用 UTF-8 或 GB18030")
    raise ValueError("SKU 成本表仅支持 CSV、XLS、XLSX")


def parse_cost_file(name: str, data: bytes) -> list[dict]:
    table = _read_table(name, data)
    header_row = sku_col = cost_col = None
    partial_headers: set[str] = set()
    for row_index, row in enumerate(table[:30]):
        normalized = [_header(value) for value in row]
        sku = next((i for i, value in enumerate(normalized) if value in SKU_HEADERS), None)
        cost = next((i for i, value in enumerate(normalized) if value in COST_HEADERS), None)
        partial_headers.update(value for value in normalized if value in PARTIAL_COST_HEADERS)
        if sku is not None and cost is not None:
            header_row, sku_col, cost_col = row_index, sku, cost
            break
    if header_row is None:
        if partial_headers:
            found = "、".join(sorted(partial_headers))
            raise ValueError(f"不支持按零散成本导入（检测到：{found}），请先汇总为“总成本”列")
        raise ValueError("找不到“SKU编码”和“总成本”两列表头")

    items: dict[str, dict] = {}
    for row_index, row in enumerate(table[header_row + 1 :], header_row + 2):
        sku = _text(row[sku_col] if sku_col < len(row) else "")
        raw_cost = row[cost_col] if cost_col < len(row) else ""
        if not sku and not _text(raw_cost):
            continue
        if _header(sku) in {"合计", "总计"}:
            continue
        if not sku:
            raise ValueError(f"第 {row_index} 行 SKU编码为空")
        if not _text(raw_cost):
            raise ValueError(f"第 {row_index} 行“{sku}”的总成本为空")
        cost = _cost_cents(raw_cost, row_index)
        if sku in items and items[sku]["unit_cost_cents"] != cost:
            raise ValueError(f"SKU编码“{sku}”在文件中重复且总成本不一致")
        items[sku] = {"sku_key": sku, "unit_cost_cents": cost}
    if not items:
        raise ValueError("成本表中没有可导入的数据")
    return list(items.values())


def save_cost_items(conn: sqlite3.Connection, shop_name: str, items: list[dict]) -> dict:
    shop_name = str(shop_name or "").strip()
    if not shop_name:
        raise ValueError("请先选择店铺，再导入或保存 SKU 成本")
    now = datetime.now().isoformat(timespec="seconds")
    counts = {"new": 0, "updated": 0, "existing": 0}
    for item in items:
        sku = _text(item.get("sku_key"))
        if not sku:
            continue
        cents = item.get("unit_cost_cents")
        if cents is None:
            cents = _cost_cents(item.get("unit_cost"), 0)
        old = conn.execute(
            "SELECT unit_cost_cents FROM sku_costs WHERE shop_name=? AND sku_key=?",
            (shop_name, sku),
        ).fetchone()
        if old is None:
            status = "new"
        elif old["unit_cost_cents"] == cents:
            status = "existing"
        else:
            status = "updated"
        conn.execute(
            """INSERT INTO sku_costs(shop_name,sku_key,unit_cost_cents,updated_at)
               VALUES(?,?,?,?) ON CONFLICT(shop_name,sku_key) DO UPDATE SET
               unit_cost_cents=excluded.unit_cost_cents,updated_at=excluded.updated_at""",
            (shop_name, sku, cents, now),
        )
        counts[status] += 1
    return counts
