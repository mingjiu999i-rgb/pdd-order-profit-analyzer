from __future__ import annotations

import sqlite3
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .engine import audit_summary, combined_summary, daily_summary

RED, DARK, LIGHT, GREEN, YELLOW, WHITE = "E02E24", "26323C", "F4F6F8", "E8F5ED", "FFF4D6", "FFFFFF"
THIN = Side(style="thin", color="E1E6E9")


def _title(ws, title, subtitle, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws["A1"] = title; ws["A1"].font = Font(size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=RED); ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws["A2"] = subtitle; ws["A2"].font = Font(size=10, color="687582"); ws.row_dimensions[2].height = 26


def _header(ws, row, headers):
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row, col, value); cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=DARK); cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _finish(ws, widths, freeze):
    ws.freeze_panes = freeze; ws.auto_filter.ref = ws.dimensions; ws.sheet_view.showGridLines = False
    for i, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows():
        for cell in row: cell.alignment = Alignment(vertical="center")


def build_xlsx(conn: sqlite3.Connection, shop_name: str | None = None, start_date: str | None = None, end_date: str | None = None) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "经营汇总"
    days = daily_summary(conn, start_date, end_date, shop_name)
    audit = audit_summary(conn, shop_name)
    total = combined_summary(days)
    as_of = datetime.now().strftime("%Y-%m-%d %H:%M")
    shops = "、".join(sorted({d["shop_name"] for d in days})) or "未识别店铺"
    headers = ["店铺名称", "成交日期", "订单数", "订单成交额", "有效订单销售额", "已发货退款", "其他扣款", "订单净回款", "回款差额", "入账完成率", "净结算率", "入账状态", "说明"]
    _title(ws, f"{shops}｜拼多多订单经营分析", f"数据生成时间：{as_of}　口径：按订单原始支付日期归集，金额单位为人民币元", len(headers))
    ws["A3"] = "说明：回款差额＝有效订单销售额－订单净回款；入账未完成时，差额包含尚未结算金额，不能全部视为费用。"
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers)); ws["A3"].font = Font(size=10, color="5C6872")
    _header(ws, 4, headers)
    report_rows = ([total] if total else []) + days
    for n, d in enumerate(report_rows, 5):
        is_total = total is not None and n == 5
        status = "入账已完成" if d["mature"] else "入账未完成"
        notes = []
        if not d["mature"]: notes.append(f"尚有 {d['eligible_orders']-d['settled_orders']} 单未出现正向结算")
        if d["unknown_funds"]: notes.append(f"{d['unknown_funds']} 笔资金类型待确认")
        if d["settlement_difference"] not in (None, 0): notes.append(f"结算差异 {d['settlement_difference']:.2f} 元")
        shop_label = f"{d['shop_name']}（累计合计）" if is_total else d["shop_name"]
        values = [shop_label, d["date"], d["order_count"], d["original_sales"], d["effective_sales"], d["shipped_refund"], d["other_deductions"], d["current_net"], d["collection_gap"], d["completion_rate"], d["net_settlement_rate"], status, "；".join(notes)]
        for c, value in enumerate(values, 1):
            cell = ws.cell(n, c, value); cell.border = Border(bottom=THIN)
            if is_total:
                cell.fill = PatternFill("solid", fgColor="E8EEF3"); cell.font = Font(bold=True, color=DARK)
            elif n % 2 == 0: cell.fill = PatternFill("solid", fgColor="FAFBFC")
        for c in range(4, 10): ws.cell(n, c).number_format = '¥#,##0.00;[Red]-¥#,##0.00'
        for c in (10, 11): ws.cell(n, c).number_format = "0.00%"
        ws.cell(n, 12).fill = PatternFill("solid", fgColor=GREEN if d["mature"] else YELLOW)
    _finish(ws, [20,13,10,14,17,14,13,15,14,13,13,13,34], "A5")
    _orders(wb, conn, shop_name, start_date, end_date); _refunds(wb, conn, shop_name, start_date, end_date); _funds(wb, conn, shop_name, start_date, end_date); _audit(wb, audit, as_of)
    out = BytesIO(); wb.save(out); return out.getvalue()


def _order_filter(shop_name=None, start_date=None, end_date=None, alias=""):
    prefix = f"{alias}." if alias else ""
    clauses, params = [], []
    if shop_name: clauses.append(f"{prefix}shop_name=?"); params.append(shop_name)
    if start_date: clauses.append(f"{prefix}pay_date>=?"); params.append(start_date)
    if end_date: clauses.append(f"{prefix}pay_date<=?"); params.append(end_date)
    return (" WHERE " + " AND ".join(clauses) if clauses else ""), params


def _orders(wb, conn, shop_name=None, start_date=None, end_date=None):
    ws = wb.create_sheet("订单明细"); headers = ["店铺名称","订单号","支付日期","支付时间","商家实收金额","订单状态","售后状态","商品ID","商品规格","数量"]
    _title(ws,"订单明细","订单号使用文本格式，保证长编号完整显示",len(headers)); _header(ws,4,headers)
    where, params = _order_filter(shop_name, start_date, end_date)
    rows = conn.execute(f"SELECT * FROM orders{where} ORDER BY pay_time,order_id", params)
    for n,r in enumerate(rows,5):
        vals=[r["shop_name"],r["order_id"],r["pay_date"],r["pay_time"],r["merchant_receipt_cents"]/100,r["order_status"],r["aftersale_status"],r["product_id"],r["sku"],r["quantity"]]
        for c,v in enumerate(vals,1): ws.cell(n,c,v).border=Border(bottom=THIN)
        ws.cell(n,2).number_format="@"; ws.cell(n,5).number_format='¥#,##0.00'
    _finish(ws,[20,27,13,21,15,18,18,16,58,9],"A5")


def _refunds(wb, conn, shop_name=None, start_date=None, end_date=None):
    ws=wb.create_sheet("退款明细"); headers=["店铺名称","售后编号","订单号","交易金额","退款金额","售后状态","发货阶段","退款类型","申请时间","同意退款时间"]
    _title(ws,"退款明细","退款表判断售后状态和发货阶段；实际资金扣款以资金明细为准",len(headers)); _header(ws,4,headers)
    where, params = _order_filter(shop_name, start_date, end_date, "o")
    rows = conn.execute(f"SELECT r.* FROM refunds r JOIN orders o ON o.order_id=r.order_id{where} ORDER BY r.apply_time,r.aftersale_id", params)
    for n,r in enumerate(rows,5):
        vals=[r["shop_name"],r["aftersale_id"],r["order_id"],r["trade_amount_cents"]/100,r["refund_amount_cents"]/100,r["aftersale_status"],r["shipping_stage"],r["refund_type"],r["apply_time"],r["agree_time"]]
        for c,v in enumerate(vals,1): ws.cell(n,c,v).border=Border(bottom=THIN)
        ws.cell(n,2).number_format="@"; ws.cell(n,3).number_format="@"; ws.cell(n,4).number_format='¥#,##0.00'; ws.cell(n,5).number_format='¥#,##0.00'
    _finish(ws,[20,20,27,13,13,15,12,12,21,21],"A5")


def _funds(wb, conn, shop_name=None, start_date=None, end_date=None):
    ws=wb.create_sheet("资金明细"); headers=["店铺名称","订单号","发生时间","收入","支出","账务类型","业务代码","业务描述","系统分类","备注"]
    _title(ws,"资金明细","每个可验证资金事件仅计算一次；全部原始导入行仍保存在本地数据库",len(headers)); _header(ws,4,headers)
    labels={"positive_settlement":"正向结算","refund":"退款","other_operating":"其他经营收支","transfer":"资金划转","promotion_spend":"推广实际消耗","unknown":"待确认"}
    where, params = _order_filter(shop_name, start_date, end_date, "o")
    rows = conn.execute(f"SELECT f.* FROM fund_events f JOIN orders o ON o.order_id=f.order_id{where} ORDER BY f.occurred_at,f.id", params)
    for n,r in enumerate(rows,5):
        vals=[r["shop_name"],r["order_id"],r["occurred_at"],r["income_cents"]/100,r["expense_cents"]/100,r["account_type"],r["business_code"],r["business_description"],labels.get(r["category"],r["category"]),r["note"]]
        for c,v in enumerate(vals,1): ws.cell(n,c,v).border=Border(bottom=THIN)
        ws.cell(n,2).number_format="@"; ws.cell(n,4).number_format='¥#,##0.00'; ws.cell(n,5).number_format='¥#,##0.00;[Red]-¥#,##0.00'
    _finish(ws,[20,27,21,12,12,16,13,34,16,18],"A5")


def _audit(wb, audit, as_of):
    ws=wb.create_sheet("导入检查"); _title(ws,"导入检查",f"检查时间：{as_of}",10)
    for row,(label,value) in enumerate((("未匹配退款",audit["unmatched_refunds"]),("未匹配资金",audit["unmatched_funds"]),("未知资金类型",audit["unknown_funds"])),4):
        ws.cell(row,1,label).font=Font(bold=True); ws.cell(row,1).fill=PatternFill("solid",fgColor=LIGHT); ws.cell(row,2,value)
    headers=["批次","店铺名称","文件名","报表类型","导入时间","原始行","新增","更新","已存在","确定重复","疑似重复"]; _header(ws,9,headers)
    labels={"order":"订单","refund":"退款","fund":"资金"}
    for n,b in enumerate(audit["batches"],10):
        vals=[b["id"],b["shop_name"],b["file_name"],labels.get(b["report_type"],b["report_type"]),b["imported_at"],b["total_rows"],b["new_rows"],b["updated_rows"],b["existing_rows"],b["duplicate_rows"],b["suspected_rows"]]
        for c,v in enumerate(vals,1): ws.cell(n,c,v).border=Border(bottom=THIN)
    _finish(ws,[9,20,55,12,21,11,10,10,11,12,12],"A10")
