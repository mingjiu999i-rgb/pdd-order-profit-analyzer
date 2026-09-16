from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sqlite3
import zipfile
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable


class ShopRequiredError(ValueError):
    """The report is valid, but its shop cannot be proven from its contents/history."""

    def __init__(self, file_name: str, report_type: str, message: str = "无法可靠识别店铺"):
        super().__init__(message)
        self.file_name = file_name
        self.report_type = report_type


def clean(value) -> str:
    return "" if value is None else str(value).replace("\t", "").strip()


def cents(value) -> int:
    text = clean(value).replace(",", "").replace("¥", "")
    if not text or text == "-":
        return 0
    try:
        return int((Decimal(text) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except InvalidOperation as exc:
        raise ValueError(f"无法识别金额：{value}") from exc


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError("CSV 编码无法识别，请使用 UTF-8 或 GB18030")


def normalize(row: dict) -> dict:
    return {clean(k): clean(v) for k, v in row.items() if k is not None}


def row_sig(row: dict, fields: list[str] | None = None) -> str:
    body = {k: row.get(k, "") for k in fields} if fields else row
    return sha256(json.dumps(body, ensure_ascii=False, sort_keys=True).encode())


def classify_fund(account_type: str, business: str) -> str:
    text = f"{account_type}|{business}"
    code = business.split("|", 1)[0].strip()
    if code in {"0010002", "0010005"}:
        return "positive_settlement"
    if code in {"0020002", "0020005"}:
        return "refund"
    if any(x in text for x in ("提现", "充值", "资金划转", "转账")):
        return "transfer"
    if any(x in text for x in ("推广消耗", "广告消耗", "推广实际扣费", "广告实际扣费", "推广费扣款", "广告费扣款")):
        return "promotion_spend"
    known_other = ("技术服务费", "售后费用", "消费者体验提升计划", "多多进宝", "小额打款", "申诉补回", "费用返还", "其他收入")
    if any(x in text for x in known_other) or code.startswith(("003", "004", "006", "013")):
        return "other_operating"
    return "unknown"


def shop_from_filename(name: str) -> str | None:
    """从商家自行命名的报表文件识别店铺名；通用平台文件名返回空。"""
    stem = Path(name).stem.strip()
    lower = stem.lower()
    if lower.startswith(("pdd-mall-bill-detail", "mall-bill-detail")) or "orders_export" in lower or re.match(r"^[0-9a-f]{20,}(?:_|$)", lower):
        return None
    text = re.sub(r"(?:截至|截止)\d{1,2}月\d{1,2}日", "", stem)
    text = re.sub(r"\d{4}[-年]?\d{1,2}月?", "", text)
    text = re.sub(r"\d{1,2}月份?", "", text)
    text = re.sub(r"(订单|退款|资金明细|资金流水|售后|推广费|推广报表|推广|报表|数据)+$", "", text).strip(" _-（）()")
    return text or None


def shop_from_content(name: str, data: bytes) -> str | None:
    """识别报表顶部常见的“店铺名称：xxx”或“店铺：xxx”字段。"""
    if Path(name).suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            if ws.calculate_dimension() == "A1:A1": ws.reset_dimensions()
            text = "\n".join(clean(v) for row in ws.iter_rows(max_row=10, values_only=True) for v in row if clean(v))
        except Exception:
            return None
    else:
        try: text = "\n".join(decode_csv(data).splitlines()[:15])
        except Exception: return None
    match = re.search(r"店铺(?:名称|名)?\s*[：:,，]\s*([^,，\n\r]+)", text)
    return clean(match.group(1)) if match else None


class Importer:
    def __init__(self, conn: sqlite3.Connection):
        self.conn = conn

    def import_path(self, path: str | Path, shop_name: str | None = None) -> list[dict]:
        path = Path(path)
        data = path.read_bytes()
        if path.suffix.lower() == ".zip":
            results = []
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                for info in archive.infolist():
                    if info.is_dir() or Path(info.filename).suffix.lower() not in {".csv", ".xls", ".xlsx"}:
                        continue
                    results.append(self.import_bytes(Path(info.filename).name, archive.read(info), shop_name))
            if not results:
                raise ValueError("ZIP 中没有 CSV 或 XLSX 报表")
            return results
        return [self.import_bytes(path.name, data, shop_name)]

    def import_bytes(self, name: str, data: bytes, shop_name: str | None = None) -> dict:
        report_type, rows, date_range = self._parse(name, data)
        shop_name = clean(shop_name) or shop_from_content(name, data) or self._detect_shop(name, report_type, rows)
        self.current_shop = shop_name
        now = datetime.now().isoformat(timespec="seconds")
        cur = self.conn.execute(
            "INSERT INTO import_batches(file_name,file_sha256,report_type,shop_name,imported_at,source_start,source_end,total_rows) VALUES(?,?,?,?,?,?,?,?)",
            (name, sha256(data), report_type, shop_name, now, date_range[0], date_range[1], len(rows)),
        )
        batch_id = cur.lastrowid
        counts = Counter()
        ordinals = Counter()
        try:
            for number, row in enumerate(rows, 1):
                signature = row_sig(row)
                ordinals[signature] += 1
                ordinal = ordinals[signature]
                raw_id = self.conn.execute(
                    "INSERT INTO raw_rows(batch_id,report_type,row_number,row_json,row_signature,event_ordinal) VALUES(?,?,?,?,?,?)",
                    (batch_id, report_type, number, json.dumps(row, ensure_ascii=False), signature, ordinal),
                ).lastrowid
                status = getattr(self, f"_save_{report_type}")(row, batch_id, raw_id, signature, ordinal, now)
                counts[status] += 1
            self.conn.execute(
                "UPDATE import_batches SET new_rows=?,updated_rows=?,existing_rows=?,duplicate_rows=?,suspected_rows=? WHERE id=?",
                (counts["new"], counts["updated"], counts["existing"], counts["duplicate"], counts["suspected"], batch_id),
            )
            self.conn.commit()
        except Exception as exc:
            self.conn.rollback()
            raise ValueError(f"{name} 第 {number} 行导入失败：{exc}") from exc
        return {"batch_id": batch_id, "file_name": name, "report_type": report_type, "shop_name": shop_name, "total": len(rows), **counts}

    def _detect_shop(self, name: str, report_type: str, rows: list[dict]) -> str:
        key = "订单号" if report_type == "order" else ("订单编号" if report_type == "refund" else "商户订单号")
        ids = list(dict.fromkeys(r.get(key) for r in rows if r.get(key)))
        if ids:
            counts = Counter()
            for start in range(0, len(ids), 500):
                chunk = ids[start:start + 500]
                marks = ",".join("?" * len(chunk))
                for match in self.conn.execute(
                    f"SELECT shop_name,COUNT(*) n FROM orders WHERE order_id IN ({marks}) AND shop_name IS NOT NULL AND shop_name<>'' GROUP BY shop_name",
                    chunk,
                ):
                    counts[match["shop_name"]] += match["n"]
            if len(counts) == 1:
                return next(iter(counts))
            if len(counts) > 1:
                detail = "、".join(f"{shop} {count}单" for shop, count in counts.most_common())
                raise ShopRequiredError(name, report_type, f"同一文件匹配到多个店铺（{detail}），已停止导入")
        # 文件名只作为新店铺首次导入的便利兜底；后续识别优先使用完整订单号。
        named = shop_from_filename(name)
        if named:
            return named
        raise ShopRequiredError(name, report_type, "无法通过报表内容或历史订单号确定店铺，请确认一次店铺名称")

    def _parse(self, name: str, data: bytes):
        suffix = Path(name).suffix.lower()
        if suffix in {".xls", ".xlsx"}:
            table = []
            if suffix == ".xlsx":
                from openpyxl import load_workbook
                wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                ws = wb["售后信息"] if "售后信息" in wb.sheetnames else wb[wb.sheetnames[0]]
                if ws.calculate_dimension() == "A1:A1": ws.reset_dimensions()
                table = [list(row) for row in ws.iter_rows(values_only=True)]
            else:
                import xlrd
                wb = xlrd.open_workbook(file_contents=data)
                ws = wb.sheet_by_index(0)
                for row_index in range(ws.nrows):
                    values = []
                    for col_index in range(ws.ncols):
                        cell = ws.cell(row_index, col_index)
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(cell.value, wb.datemode).strftime("%Y-%m-%d %H:%M:%S")
                        else:
                            value = cell.value
                        values.append(value)
                    table.append(values)
            header_idx = next((i for i, row in enumerate(table[:30]) if self._is_supported_header({clean(x) for x in row})), None)
            if header_idx is None: raise ValueError("找不到退款或推广报表表头")
            headers = [clean(x) for x in table[header_idx]]
            rows = [normalize(dict(zip(headers, row))) for row in table[header_idx + 1:] if any(clean(x) for x in row)]
            if "售后编号" in headers and "订单编号" in headers:
                dates = [r.get("申请时间", "") for r in rows if r.get("申请时间")]
                return "refund", rows, (min(dates, default=None), max(dates, default=None))
            promotion = self._prepare_promotion(rows, set(headers))
            if promotion: return promotion
            raise ValueError("Excel 不是支持的退款或推广报表")
        text = decode_csv(data)
        lines = text.splitlines()
        header_idx = next((i for i, line in enumerate(lines[:15]) if ("订单号" in line or "商户订单号" in line) and "," in line), None)
        if header_idx is None:
            raise ValueError("找不到 CSV 表头")
        rows = [normalize(r) for r in csv.DictReader(io.StringIO("\n".join(lines[header_idx:]))) if any(clean(v) for v in r.values())]
        headers = set(rows[0]) if rows else set()
        if "订单号" in headers and ("支付时间" in headers or "订单成交时间" in headers):
            # 拼多多新版订单导出把原“支付时间”改名为“订单成交时间”。
            # 两种模板统一映射到内部支付时间，保证历史和新版报表可混合导入。
            if "支付时间" not in headers:
                for row in rows:
                    row["支付时间"] = row.get("订单成交时间", "")
            dates = [r.get("支付时间", "") for r in rows if r.get("支付时间")]
            return "order", rows, (min(dates, default=None), max(dates, default=None))
        if "商户订单号" in headers and "发生时间" in headers:
            rows = [r for r in rows if re.fullmatch(r"\d{6}-\d+", r.get("商户订单号", ""))]
            dates = [r.get("发生时间", "") for r in rows if r.get("发生时间")]
            return "fund", rows, (min(dates, default=None), max(dates, default=None))
        promotion = self._prepare_promotion(rows, headers)
        if promotion: return promotion
        raise ValueError("CSV 不是支持的订单或资金报表")

    @staticmethod
    def _promotion_columns(headers):
        date_col = next((x for x in ("统计日期", "日期", "消耗日期", "推广日期", "数据日期") if x in headers), None)
        amount_names = ("消耗金额(元)", "消耗金额（元）", "推广消耗(元)", "推广消耗（元）", "花费(元)", "花费（元）", "实际消耗(元)", "实际消耗（元）", "总花费(元)", "总花费（元）", "消耗", "花费")
        amount_col = next((x for x in amount_names if x in headers), None)
        return date_col, amount_col

    @classmethod
    def _is_supported_header(cls, headers):
        return ({"售后编号", "订单编号"} <= headers) or all(cls._promotion_columns(headers))

    @classmethod
    def _prepare_promotion(cls, rows, headers):
        date_col, amount_col = cls._promotion_columns(headers)
        if not date_col or not amount_col: return None
        valid_rows = []
        for row in rows:
            text = clean(row.get(date_col, "")).replace("/", "-").replace("年", "-").replace("月", "-").replace("日", "")
            match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:\s.*)?", text)
            if not match: continue  # 跳过“总计”等非日期行。
            row["__promotion_date"] = f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
            row["__promotion_amount"] = row.get(amount_col, "")
            if clean(row["__promotion_amount"]): valid_rows.append(row)
        rows[:] = valid_rows
        dates = [r["__promotion_date"] for r in rows]
        return "promotion", rows, (min(dates, default=None), max(dates, default=None))

    def _save_order(self, r, batch, raw, sig, ordinal, now):
        order_id = r.get("订单号", "")
        if not order_id:
            raise ValueError("订单号为空")
        pay_time = r.get("支付时间") or None
        pay_date = pay_time[:10] if pay_time else None
        values = (pay_time, pay_date, self.current_shop, cents(r.get("商家实收金额(元)")), r.get("订单状态"), r.get("售后状态"), r.get("商品id"), r.get("商品规格"), clean(r.get("商家编码-规格维度")), int(r.get("商品数量(件)") or 0), batch, now, order_id)
        old = self.conn.execute("SELECT * FROM orders WHERE order_id=?", (order_id,)).fetchone()
        if not old:
            self.conn.execute("INSERT INTO orders(pay_time,pay_date,shop_name,merchant_receipt_cents,order_status,aftersale_status,product_id,sku,sku_code,quantity,source_batch_id,updated_at,order_id) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)", values)
            return "new"
        comparable = tuple(old[k] for k in ("pay_time","pay_date","shop_name","merchant_receipt_cents","order_status","aftersale_status","product_id","sku","sku_code","quantity"))
        if comparable == values[:10]:
            return "existing"
        self.conn.execute("UPDATE orders SET pay_time=?,pay_date=?,shop_name=?,merchant_receipt_cents=?,order_status=?,aftersale_status=?,product_id=?,sku=?,sku_code=?,quantity=?,source_batch_id=?,updated_at=? WHERE order_id=?", values)
        return "updated"

    def _save_refund(self, r, batch, raw, sig, ordinal, now):
        aid = r.get("售后编号", "")
        if not aid:
            raise ValueError("售后编号为空")
        values = (r.get("订单编号"), self.current_shop, cents(r.get("交易金额")), cents(r.get("退款金额")), r.get("售后状态"), r.get("订单状态"), r.get("退款类型"), r.get("申请时间"), r.get("同意退款时间"), batch, now, aid)
        old = self.conn.execute("SELECT * FROM refunds WHERE aftersale_id=?", (aid,)).fetchone()
        if not old:
            self.conn.execute("INSERT INTO refunds(order_id,shop_name,trade_amount_cents,refund_amount_cents,aftersale_status,shipping_stage,refund_type,apply_time,agree_time,source_batch_id,updated_at,aftersale_id) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", values)
            return "new"
        comparable = tuple(old[k] for k in ("order_id","shop_name","trade_amount_cents","refund_amount_cents","aftersale_status","shipping_stage","refund_type","apply_time","agree_time"))
        if comparable == values[:9]:
            return "existing"
        self.conn.execute("UPDATE refunds SET order_id=?,shop_name=?,trade_amount_cents=?,refund_amount_cents=?,aftersale_status=?,shipping_stage=?,refund_type=?,apply_time=?,agree_time=?,source_batch_id=?,updated_at=? WHERE aftersale_id=?", values)
        return "updated"

    def _save_fund(self, r, batch, raw, sig, ordinal, now):
        fields = ["商户订单号","发生时间","收入金额（+元）","支出金额（-元）","账务类型","备注","业务描述"]
        canonical_sig = row_sig(r, fields)
        key = f"{canonical_sig}:{ordinal}"
        old = self.conn.execute("SELECT id FROM fund_events WHERE event_key=?", (key,)).fetchone()
        if old:
            self.conn.execute("UPDATE fund_events SET shop_name=? WHERE id=? AND (shop_name IS NULL OR shop_name='' OR shop_name='未识别店铺')", (self.current_shop, old["id"]))
            self.conn.execute("INSERT INTO fund_event_sources(fund_event_id,raw_row_id) VALUES(?,?)", (old["id"], raw))
            return "duplicate"
        business = r.get("业务描述", "")
        code, _, desc = business.partition("|")
        category = classify_fund(r.get("账务类型", ""), business)
        cur = self.conn.execute(
            "INSERT INTO fund_events(event_key,row_signature,event_ordinal,order_id,shop_name,occurred_at,income_cents,expense_cents,account_type,business_code,business_description,note,category,source_batch_id) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (key, canonical_sig, ordinal, r.get("商户订单号"), self.current_shop, r.get("发生时间"), cents(r.get("收入金额（+元）")), cents(r.get("支出金额（-元）")), r.get("账务类型"), code, desc, r.get("备注"), category, batch),
        )
        self.conn.execute("INSERT INTO fund_event_sources(fund_event_id,raw_row_id) VALUES(?,?)", (cur.lastrowid, raw))
        return "new"

    def _save_promotion(self, r, batch, raw, sig, ordinal, now):
        spend_date = clean(r.get("__promotion_date"))[:10]
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", spend_date):
            raise ValueError("推广报表的统计日期格式不正确")
        amount = abs(cents(r.get("__promotion_amount")))
        key = f"{sig}:{ordinal}"
        if self.conn.execute("SELECT id FROM promotion_expenses WHERE row_key=?", (key,)).fetchone():
            return "duplicate"
        self.conn.execute(
            "INSERT INTO promotion_expenses(row_key,shop_name,spend_date,amount_cents,source_batch_id) VALUES(?,?,?,?,?)",
            (key, self.current_shop, spend_date, amount, batch),
        )
        return "new"
